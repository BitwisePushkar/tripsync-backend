import io
import logging
from django.utils import timezone
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from .models import Itinerary 

logger = logging.getLogger(__name__)

def _fmt_cost(val):
    try:
        return f"₹{float(val):,.0f}"
    except Exception:
        return str(val)

def _time_range(activity):
    if activity.start_time and activity.end_time:
        return f"{activity.start_time.strftime('%H:%M')} – {activity.end_time.strftime('%H:%M')}"
    return activity.get_time_display().capitalize()

def generate_pdf(trip) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm,
        topMargin=2*cm, bottomMargin=2*cm,
    )

    styles = getSampleStyleSheet()
    BRAND = colors.HexColor("#2563EB")
    LIGHT = colors.HexColor("#EFF6FF")
    GRAY = colors.HexColor("#6B7280")
    DANGER = colors.HexColor("#DC2626")
    h1 = ParagraphStyle("H1", parent=styles["Title"],  fontSize=22, textColor=BRAND, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=14, textColor=BRAND, spaceBefore=14, spaceAfter=4)
    h3 = ParagraphStyle("H3", parent=styles["Heading3"], fontSize=11, textColor=colors.black, spaceBefore=8, spaceAfter=2)
    normal = ParagraphStyle("N", parent=styles["Normal"], fontSize=9, leading=13)
    small = ParagraphStyle("S", parent=styles["Normal"], fontSize=8, textColor=GRAY)
    warn = ParagraphStyle("W", parent=styles["Normal"], fontSize=8, textColor=DANGER)
    story = []

    story.append(Paragraph(trip.tripname, h1))
    story.append(Paragraph(
        f"{trip.current_loc} → <b>{trip.destination}</b> &nbsp;|&nbsp; "
        f"{trip.start_date} – {trip.end_date} ({trip.days} days) &nbsp;|&nbsp; "
        f"{trip.get_trip_type_display()} · {trip.get_trip_preferences_display()}",
        normal,
    ))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1, color=BRAND))
    story.append(Spacer(1, 8))

    story.append(Paragraph("Budget Overview", h2))
    bdata = [["Category", "Allocated", "Spent", "Remaining"]]
    for bc in trip.budget_categories.all():
        bdata.append([
            bc.get_category_display(),
            _fmt_cost(bc.allocated),
            _fmt_cost(bc.spent),
            _fmt_cost(bc.remaining),
        ])
    bdata.append(["TOTAL", _fmt_cost(trip.total_budget), "", _fmt_cost(trip.budget_remaining)])
    bt = Table(bdata, colWidths=[5*cm, 3.5*cm, 3.5*cm, 3.5*cm])
    bt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTSIZE",   (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, -1), (-1, -1), LIGHT),
        ("FONTNAME",   (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F9FAFB")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(bt)
    story.append(Spacer(1, 12))

    try:
        profile = trip.user.profile
        personal_contacts = list(profile.emergency_contacts.all())
    except Exception:
        personal_contacts = []

    if personal_contacts:
        story.append(Paragraph("Personal Emergency Contacts", h2))
        pcdata = [["Name", "Relation", "Phone Number", "Email"]]
        for pc in personal_contacts:
            pcdata.append([
                pc.name,
                pc.relation,
                pc.phone_number,
                pc.email,
            ])
        pct = Table(pcdata, colWidths=[4.2*cm, 2.8*cm, 4.2*cm, 4.3*cm])
        pct.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), DANGER),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTSIZE",   (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF5F5")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#FCA5A5")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(pct)
        story.append(Spacer(1, 12))

    try:
        itinerary = trip.itinerary
        for day in itinerary.day_plans.all():
            story.append(Paragraph(f"Day {day.day_number} — {day.title}", h2))
            if day.date:
                story.append(Paragraph(str(day.date.strftime("%A, %d %b %Y")), small))
            if day.tips:
                story.append(Paragraph(f"💡 {day.tips}", small))
            story.append(Spacer(1, 4))

            for act in day.activities.ordered_by_time():
                story.append(Paragraph(f"🕐 {_time_range(act)}  {act.title}", h3))
                story.append(Paragraph(f"📍 {act.location_name or '—'}", small))
                if act.place and act.place.address:
                    story.append(Paragraph(f"Address: {act.place.address}", small))
                if act.place and act.place.phone:
                    story.append(Paragraph(f"📞 {act.place.phone}", small))
                if act.place and act.place.rating:
                    story.append(Paragraph(f"⭐ Rating: {act.place.rating}", small))
                story.append(Paragraph(
                    f"Category: {act.get_category_display()}  |  Est. cost: {_fmt_cost(act.estimated_cost)}"
                    + (f"  |  Distance: {act.distance_km} km" if act.distance_km else ""),
                    small,
                ))
                if act.short_description:
                    story.append(Paragraph(act.short_description, normal))
                if act.dos_and_donts:
                    story.append(Paragraph(f"<i>{act.dos_and_donts}</i>", small))
                if act.emergency_contacts and isinstance(act.emergency_contacts, list):
                    contacts_list = []
                    for c in act.emergency_contacts:
                        if isinstance(c, dict):
                            contacts_list.append(f"{c.get('name', '?')} {c.get('phone', '')}".strip())
                        elif isinstance(c, str):
                            contacts_list.append(c)
                    if contacts_list:
                        contacts = "  |  ".join(contacts_list)
                        story.append(Paragraph(f"🚨 Emergency: {contacts}", warn))

                story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E5E7EB")))
    except Itinerary.DoesNotExist:
        story.append(Paragraph("No itinerary generated yet.", normal))

    story.append(Spacer(1, 16))
    story.append(Paragraph(
        f"Generated by TripX  ·  {timezone.now().strftime('%d %b %Y %H:%M')}",
        small,
    ))

    doc.build(story)
    return buf.getvalue()

def generate_docx(trip) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    BRAND_RGB = RGBColor(0x25, 0x63, 0xEB)
    GRAY_RGB  = RGBColor(0x6B, 0x72, 0x80)

    document = Document()

    for section in document.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    def add_heading(text, level=1, color=BRAND_RGB):
        p = document.add_heading(text, level=level)
        for run in p.runs:
            run.font.color.rgb = color
        return p

    def add_para(text, size=9, color=None, italic=False, bold=False):
        p = document.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        run.font.italic = italic
        run.font.bold   = bold
        return p

    def add_hr():
        p = document.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "2563EB")
        pBdr.append(bottom)
        pPr.append(pBdr)

    add_heading(trip.tripname, 1)
    add_para(
        f"{trip.current_loc} → {trip.destination}  |  "
        f"{trip.start_date} – {trip.end_date} ({trip.days} days)  |  "
        f"{trip.get_trip_type_display()} · {trip.get_trip_preferences_display()}",
        size=10,
    )
    add_hr()

    add_heading("Budget Overview", 2)
    headers = ["Category", "Allocated", "Spent", "Remaining"]
    rows = []
    for bc in trip.budget_categories.all():
        rows.append([
            bc.get_category_display(),
            _fmt_cost(bc.allocated),
            _fmt_cost(bc.spent),
            _fmt_cost(bc.remaining),
        ])
    rows.append(["TOTAL", _fmt_cost(trip.total_budget), "", _fmt_cost(trip.budget_remaining)])

    tbl = document.add_table(rows=1 + len(rows), cols=4)
    tbl.style = "Table Grid"
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.font.bold  = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "2563EB")
        tcPr.append(shd)

    for ri, row_data in enumerate(rows):
        row = tbl.rows[ri + 1]
        for ci, val in enumerate(row_data):
            row.cells[ci].text = val

    document.add_paragraph()

    try:
        profile = trip.user.profile
        personal_contacts = list(profile.emergency_contacts.all())
    except Exception:
        personal_contacts = []

    if personal_contacts:
        add_heading("Personal Emergency Contacts", 2, color=RGBColor(0xDC, 0x26, 0x26))
        headers_pc = ["Name", "Relation", "Phone Number", "Email"]
        tbl_pc = document.add_table(rows=1 + len(personal_contacts), cols=4)
        tbl_pc.style = "Table Grid"
        hdr_cells_pc = tbl_pc.rows[0].cells
        for i, h in enumerate(headers_pc):
            hdr_cells_pc[i].text = h
            run = hdr_cells_pc[i].paragraphs[0].runs[0]
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            tc = hdr_cells_pc[i]._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "DC2626")
            tcPr.append(shd)

        for ri, pc in enumerate(personal_contacts):
            row = tbl_pc.rows[ri + 1]
            row.cells[0].text = pc.name
            row.cells[1].text = pc.relation
            row.cells[2].text = pc.phone_number
            row.cells[3].text = pc.email

        document.add_paragraph()

    try:
        itinerary = trip.itinerary
        for day in itinerary.day_plans.all():
            add_heading(f"Day {day.day_number} — {day.title}", 2)
            if day.date:
                add_para(day.date.strftime("%A, %d %b %Y"), size=9, color=GRAY_RGB)
            if day.tips:
                add_para(f"💡 {day.tips}", size=9, color=GRAY_RGB, italic=True)

            for act in day.activities.ordered_by_time():
                add_heading(f"{_time_range(act)}  {act.title}", 3)
                add_para(f"📍 {act.location_name or '—'}  |  {act.get_category_display()}  |  {_fmt_cost(act.estimated_cost)}", size=9)
                if act.place:
                    p = act.place
                    if p.address:
                        add_para(f"Address: {p.address}", size=8, color=GRAY_RGB)
                    if p.phone:
                        add_para(f"📞 {p.phone}", size=8, color=GRAY_RGB)
                    if p.rating:
                        add_para(f"⭐ {p.rating}  |  {p.website or ''}", size=8, color=GRAY_RGB)
                if act.distance_km:
                    add_para(f"Distance from previous: {act.distance_km} km", size=8, color=GRAY_RGB)
                if act.short_description:
                    add_para(act.short_description, size=9)
                if act.dos_and_donts:
                    add_para(act.dos_and_donts, size=8, color=GRAY_RGB, italic=True)
                if act.emergency_contacts and isinstance(act.emergency_contacts, list):
                    contacts_list = []
                    for c in act.emergency_contacts:
                        if isinstance(c, dict):
                            contacts_list.append(f"{c.get('name', '?')} {c.get('phone', '')}".strip())
                        elif isinstance(c, str):
                            contacts_list.append(c)
                    if contacts_list:
                        contacts = "  |  ".join(contacts_list)
                        add_para(f"🚨 Emergency contacts: {contacts}", size=8,
                                 color=RGBColor(0xDC, 0x26, 0x26))
            add_hr()
    except Exception:
        add_para("No itinerary generated yet.")
    add_para(f"Generated by Itinerary  ·  {timezone.now().strftime('%d %b %Y %H:%M')}",
             size=8, color=GRAY_RGB)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()